"""
OCR, conversion & optimization feature router.

Endpoints (all under /api/pdf):

  GET  /ocr/languages                 installed tesseract languages
  GET  /{doc_id}/ocr/detect           per-page scan detection
  POST /{doc_id}/ocr                  start an OCR job (returns job_id)
  GET  /ocr/jobs/{job_id}             poll OCR job progress / result
  GET  /{doc_id}/export/{fmt}         docx | txt | md | html | png | jpg | xlsx | csv
  POST /create                        images / text / markdown / docx / pdf -> new document
  POST /{doc_id}/compress             optimize with a preset, report before/after sizes

Coordinates: everything here is in PDF points with PyMuPDF's top-left origin
(the same space ``page.get_text("words")`` returns).  Nothing in this module
accepts screen pixels.

Every mutating endpoint calls ``snapshot(doc_id, operation)`` before it writes,
so the shared undo/redo stack works.
"""

from __future__ import annotations

import asyncio
import html as _html
import io
import os
import re
import shutil
import statistics
import tempfile
import threading
import time
import uuid
import zipfile
from pathlib import Path
from typing import Literal, Optional

import fitz  # PyMuPDF
from fastapi import APIRouter, Body, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from backend import advanced_ops
from backend.advanced_ops import snapshot

router = APIRouter(prefix="/api/pdf")

_UUID_RE = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}$")
MAX_CREATE_BYTES = int(os.environ.get("MAX_FILE_SIZE_MB", "50")) * 1024 * 1024

# Per-document locks so an OCR job and a compress call never write the same
# file at the same time.
_doc_locks: dict[str, threading.Lock] = {}
_doc_locks_guard = threading.Lock()


def _lock_for(doc_id: str) -> threading.Lock:
    with _doc_locks_guard:
        lk = _doc_locks.get(doc_id)
        if lk is None:
            lk = _doc_locks[doc_id] = threading.Lock()
        return lk


def _upload_dir() -> Path:
    # Read at call time so it always matches the directory snapshot() uses.
    return Path(advanced_ops.UPLOAD_DIR)


def _doc_path(doc_id: str) -> Path:
    if not _UUID_RE.match(doc_id):
        raise HTTPException(status_code=400, detail="Invalid document ID")
    p = _upload_dir() / doc_id / "original.pdf"
    if not p.exists():
        raise HTTPException(status_code=404, detail="Document not found")
    return p


def _open(doc_id: str) -> tuple[fitz.Document, Path]:
    p = _doc_path(doc_id)
    try:
        return fitz.open(str(p)), p
    except Exception as e:  # pragma: no cover - corrupt file
        raise HTTPException(status_code=422, detail=f"Cannot open PDF: {e}")


def _atomic_save(doc: fitz.Document, path: Path, **kw) -> None:
    tmp = str(path) + ".convert.tmp"
    doc.save(tmp, **kw)
    os.replace(tmp, str(path))


def _download(content: bytes, media_type: str, filename: str) -> Response:
    return Response(
        content=content,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


# ═════════════════════════════════════════════════════════════════════════════
#  OCR
# ═════════════════════════════════════════════════════════════════════════════

_LANG_RE = re.compile(r"^[A-Za-z_]{3,20}(\+[A-Za-z_]{3,20}){0,4}$")
_CJK_LANGS = {"chi_sim", "chi_tra", "jpn", "kor", "chi_sim_vert", "chi_tra_vert", "jpn_vert", "kor_vert"}


def _tessdata() -> Optional[str]:
    env = os.environ.get("TESSDATA_PREFIX")
    if env and Path(env).is_dir():
        return env
    try:
        td = fitz.get_tessdata()
        if td and Path(td).is_dir():
            return td
    except Exception:
        pass
    for cand in ("/opt/homebrew/share/tessdata", "/usr/local/share/tessdata",
                 "/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata"):
        if Path(cand).is_dir():
            return cand
    return None


def _available_languages() -> list[str]:
    td = _tessdata()
    if not td:
        return []
    return sorted(p.stem for p in Path(td).glob("*.traineddata") if p.stem not in ("osd", "snum"))


def _validate_language(language: str) -> str:
    if not _LANG_RE.match(language):
        raise HTTPException(status_code=400, detail="Invalid language code")
    avail = set(_available_languages())
    missing = [l for l in language.split("+") if l not in avail]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Tesseract language(s) not installed: {', '.join(missing)}. "
                   f"Installed: {', '.join(sorted(avail)) or 'none'}",
        )
    return language


def _page_scan_info(page: fitz.Page) -> dict:
    """Classify a page: visible text chars, invisible (OCR) chars, image coverage."""
    visible = 0
    invisible = 0
    try:
        for span in page.get_texttrace():
            n = sum(1 for c in span["chars"] if chr(c[0]).strip())
            if span.get("type") == 3:
                invisible += n
            else:
                visible += n
    except Exception:
        visible = len(page.get_text().strip())

    page_area = abs(page.rect) or 1.0
    covered = fitz.Rect()
    img_area = 0.0
    try:
        for info in page.get_image_info():
            r = fitz.Rect(info["bbox"]) & page.rect
            if not r.is_empty:
                img_area += abs(r)
                covered |= r
    except Exception:
        pass
    coverage = min(1.0, img_area / page_area)
    is_scanned = coverage >= 0.5 and visible < 50
    return {
        "page": page.number,
        "text_chars": visible,
        "ocr_text_chars": invisible,
        "image_coverage": round(coverage, 3),
        "is_scanned": is_scanned,
        "has_ocr_layer": invisible > 0,
        "needs_ocr": is_scanned and invisible == 0,
    }


@router.get("/ocr/languages")
async def ocr_languages():
    return {"languages": _available_languages(), "tessdata": bool(_tessdata())}


@router.get("/{doc_id}/ocr/detect")
async def ocr_detect(doc_id: str):
    doc, _ = _open(doc_id)
    try:
        pages = [_page_scan_info(p) for p in doc]
    finally:
        doc.close()
    return {
        "pages": pages,
        "scanned_pages": [p["page"] for p in pages if p["is_scanned"]],
        "needs_ocr": [p["page"] for p in pages if p["needs_ocr"]],
    }


def _sample_colors(pix: fitz.Pixmap, rect: fitz.Rect, scale: float) -> tuple[tuple, tuple]:
    """Return (background_rgb, foreground_rgb) in 0..1 floats for a word box.

    Background = median of the box border pixels, foreground = mean of the
    darkest-contrast 15% of pixels inside the box.
    """
    x0 = max(0, int(rect.x0 * scale)); y0 = max(0, int(rect.y0 * scale))
    x1 = min(pix.width - 1, int(rect.x1 * scale)); y1 = min(pix.height - 1, int(rect.y1 * scale))
    if x1 <= x0 or y1 <= y0:
        return (1.0, 1.0, 1.0), (0.0, 0.0, 0.0)

    def px(x, y):
        v = pix.pixel(x, y)
        return v[:3] if len(v) >= 3 else (v[0], v[0], v[0])

    border = []
    stepx = max(1, (x1 - x0) // 20); stepy = max(1, (y1 - y0) // 8)
    for x in range(x0, x1 + 1, stepx):
        border.append(px(x, y0)); border.append(px(x, y1))
    for y in range(y0, y1 + 1, stepy):
        border.append(px(x0, y)); border.append(px(x1, y))
    bg = tuple(statistics.median(c[i] for c in border) for i in range(3))
    inner = [px(x, y) for x in range(x0, x1 + 1, stepx) for y in range(y0, y1 + 1, stepy)]
    inner.sort(key=lambda c: -sum(abs(c[i] - bg[i]) for i in range(3)))
    top = inner[: max(1, len(inner) * 15 // 100)]
    fg = tuple(sum(c[i] for c in top) / len(top) for i in range(3))
    return tuple(v / 255 for v in bg), tuple(v / 255 for v in fg)


def _ocr_page(page: fitz.Page, language: str, dpi: int, mode: str, tessdata: Optional[str]) -> int:
    """OCR one page in place and add a text layer aligned to the scan.

    mode="searchable": invisible text (render mode 3) over the untouched image —
                       Acrobat's "Recognize Text / Searchable Image".
    mode="editable":   each recognized word is covered with its sampled
                       background colour and redrawn as real visible text in the
                       sampled ink colour, so the text editor can change it.
    Returns the number of words added.
    """
    rotation = page.rotation
    if rotation:
        page.set_rotation(0)  # work in unrotated space; restored below
    try:
        kw = {"full": True, "language": language, "dpi": dpi}
        if tessdata:
            kw["tessdata"] = tessdata
        tp = page.get_textpage_ocr(**kw)
        words = page.get_text("words", textpage=tp)
        if not words:
            return 0

        # Do not duplicate text the page already has natively.
        native = [fitz.Rect(w[:4]) for w in page.get_text("words")]

        use_cjk = any(l in _CJK_LANGS for l in language.split("+"))
        fontname = "china-s" if use_cjk else "helv"
        font = fitz.Font("cjk") if use_cjk else fitz.Font("helv")
        asc, desc = font.ascender, font.descender
        span_h = (asc - desc) or 1.0

        pix = None
        scale = 1.0
        if mode == "editable":
            scale = 150 / 72
            pix = page.get_pixmap(dpi=150, colorspace=fitz.csRGB, alpha=False)

        shape = page.new_shape()
        added = 0
        for w in words:
            text = w[4]
            if not text.strip():
                continue
            r = fitz.Rect(w[:4])
            if r.is_empty or r.height < 1 or r.width < 0.5:
                continue
            center = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
            if any(center in nr for nr in native):
                continue
            fs = max(1.0, r.height / span_h)
            baseline = fitz.Point(r.x0, r.y1 + desc * fs)
            natural = font.text_length(text, fontsize=fs) or 1.0
            sx = max(0.2, min(5.0, r.width / natural))
            morph = (baseline, fitz.Matrix(sx, 1))
            if mode == "editable":
                bg, fg = _sample_colors(pix, r, scale)
                shape.draw_rect(r + (-0.5, -0.5, 0.5, 0.5))
                shape.finish(color=None, fill=bg, width=0)
                shape.insert_text(baseline, text, fontsize=fs, fontname=fontname,
                                  color=fg, render_mode=0, morph=morph)
            else:
                shape.insert_text(baseline, text, fontsize=fs, fontname=fontname,
                                  render_mode=3, morph=morph)
            added += 1
        if added:
            shape.commit(overlay=True)
        return added
    finally:
        if rotation:
            page.set_rotation(rotation)


class OCRRequest(BaseModel):
    language: str = "eng"
    pages: Optional[list[int]] = None  # None = auto (pages that need OCR)
    dpi: int = Field(300, ge=72, le=600)
    mode: Literal["searchable", "editable"] = "searchable"
    force: bool = False  # OCR even if the page already has text / an OCR layer


_jobs: dict[str, dict] = {}
_jobs_guard = threading.Lock()


def _set_job(job_id: str, **kw):
    with _jobs_guard:
        _jobs[job_id].update(kw)


def run_ocr(doc_id: str, req: OCRRequest, job_id: Optional[str] = None) -> dict:
    """Run OCR synchronously. Used by the job thread (and directly by tests)."""
    lock = _lock_for(doc_id)
    with lock:
        path = _doc_path(doc_id)
        doc = fitz.open(str(path))
        try:
            infos = [_page_scan_info(p) for p in doc]
            if req.pages is None:
                targets = [i["page"] for i in infos if (i["is_scanned"] if req.force else i["needs_ocr"])]
            else:
                bad = [p for p in req.pages if p < 0 or p >= len(doc)]
                if bad:
                    raise HTTPException(status_code=400, detail=f"Invalid page(s): {bad}")
                targets = sorted(set(req.pages))
                if not req.force:
                    targets = [p for p in targets if not infos[p]["has_ocr_layer"]]

            result = {"pages_processed": [], "words_added": 0, "skipped": [], "total": len(targets)}
            if job_id:
                _set_job(job_id, total=len(targets), done=0, status="running")
            if not targets:
                return result

            snapshot(doc_id, f"OCR ({req.mode}, {req.language})")
            td = _tessdata()
            for n, pno in enumerate(targets):
                added = _ocr_page(doc[pno], req.language, req.dpi, req.mode, td)
                result["words_added"] += added
                (result["pages_processed"] if added else result["skipped"]).append(pno)
                if job_id:
                    _set_job(job_id, done=n + 1, current_page=pno)
            _atomic_save(doc, path, garbage=3, deflate=True)
            return result
        finally:
            doc.close()


def _job_runner(job_id: str, doc_id: str, req: OCRRequest):
    try:
        res = run_ocr(doc_id, req, job_id)
        _set_job(job_id, status="done", result=res, finished=time.time())
    except HTTPException as e:
        _set_job(job_id, status="error", error=str(e.detail), finished=time.time())
    except Exception as e:  # pragma: no cover - tesseract failures etc.
        _set_job(job_id, status="error", error=str(e), finished=time.time())


@router.post("/{doc_id}/ocr")
async def start_ocr(doc_id: str, req: OCRRequest = Body(default_factory=OCRRequest)):
    _doc_path(doc_id)
    _validate_language(req.language)
    if not _tessdata():
        raise HTTPException(status_code=503, detail="Tesseract language data not found on server")
    job_id = uuid.uuid4().hex
    with _jobs_guard:
        # drop finished jobs older than an hour
        now = time.time()
        for k in [k for k, v in _jobs.items() if v.get("finished") and now - v["finished"] > 3600]:
            del _jobs[k]
        _jobs[job_id] = {"job_id": job_id, "doc_id": doc_id, "status": "queued",
                         "done": 0, "total": None, "created": now}
    threading.Thread(target=_job_runner, args=(job_id, doc_id, req), daemon=True).start()
    return {"job_id": job_id}


@router.get("/ocr/jobs/{job_id}")
async def ocr_job(job_id: str):
    with _jobs_guard:
        job = _jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        out = dict(job)
    total = out.get("total")
    out["progress"] = 1.0 if out["status"] == "done" else (
        (out.get("done", 0) / total) if total else 0.0)
    return out


# ═════════════════════════════════════════════════════════════════════════════
#  Export
# ═════════════════════════════════════════════════════════════════════════════

def _page_text(page: fitz.Page) -> str:
    return page.get_text("text", sort=True)


def export_txt(doc: fitz.Document) -> bytes:
    # Form feed between pages, the pdftotext convention.
    return "\f".join(_page_text(p).rstrip() + "\n" for p in doc).encode("utf-8")


def _md_escape(s: str) -> str:
    return re.sub(r"([\\`*_\[\]])", r"\\\1", s)


_BULLET_RE = re.compile(r"^\s*([•●▪–\-\*·])\s+")
_NUM_RE = re.compile(r"^\s*(\d{1,3})[.)]\s+")


def _table_to_md(rows: list[list]) -> str:
    rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "\\|").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def _find_tables(page: fitz.Page) -> list:
    try:
        return list(page.find_tables().tables)
    except Exception:
        return []


def export_markdown(doc: fitz.Document) -> str:
    # Body font size = most common size by character count across the doc.
    sizes: dict[float, int] = {}
    for page in doc:
        for b in page.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    k = round(s["size"], 1)
                    sizes[k] = sizes.get(k, 0) + len(s["text"].strip())
    body = max(sizes, key=sizes.get) if sizes else 11.0

    parts: list[str] = []
    for page in doc:
        tables = _find_tables(page)
        tbl_rects = [fitz.Rect(t.bbox) for t in tables]
        items: list[tuple[float, float, str]] = []  # (y, x, markdown)
        for t in tables:
            md = _table_to_md(t.extract())
            if md:
                items.append((t.bbox[1], t.bbox[0], md))

        for b in page.get_text("dict", sort=True)["blocks"]:
            if b.get("type") != 0:
                continue
            br = fitz.Rect(b["bbox"])
            if any(abs(br & tr) > 0.5 * abs(br) for tr in tbl_rects if abs(br)):
                continue
            lines_md = []
            max_size = 0.0
            for l in b["lines"]:
                segs = []
                for s in l["spans"]:
                    txt = s["text"]
                    if not txt.strip():
                        segs.append(txt)
                        continue
                    max_size = max(max_size, s["size"])
                    esc = _md_escape(txt.strip())
                    lead = " " if txt[:1].isspace() else ""
                    trail = " " if txt[-1:].isspace() else ""
                    bold = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                    ital = bool(s["flags"] & 2) or "italic" in s["font"].lower() or "oblique" in s["font"].lower()
                    mono = bool(s["flags"] & 8) or "mono" in s["font"].lower() or "courier" in s["font"].lower()
                    if mono:
                        esc = f"`{txt.strip()}`"
                    elif bold and ital:
                        esc = f"***{esc}***"
                    elif bold:
                        esc = f"**{esc}**"
                    elif ital:
                        esc = f"*{esc}*"
                    segs.append(lead + esc + trail)
                line = re.sub(r"\s+", " ", "".join(segs)).strip()
                if line:
                    lines_md.append(line)
            if not lines_md:
                continue
            ratio = max_size / body if body else 1
            plain = re.sub(r"[*`]", "", " ".join(lines_md))
            if ratio >= 1.15 and len(plain) < 200:
                level = 1 if ratio >= 1.6 else 2 if ratio >= 1.3 else 3
                heading = re.sub(r"^\*+|\*+$", "", " ".join(lines_md)).strip()
                items.append((b["bbox"][1], b["bbox"][0], "#" * level + " " + heading))
                continue
            out_lines = []
            para: list[str] = []
            for ln in lines_md:
                raw = re.sub(r"^[*`]+", "", ln)
                m_b = _BULLET_RE.match(raw)
                m_n = _NUM_RE.match(raw)
                if m_b or m_n:
                    if para:
                        out_lines.append(" ".join(para)); para = []
                    if m_b:
                        out_lines.append("- " + _BULLET_RE.sub("", raw, count=1))
                    else:
                        out_lines.append(f"{m_n.group(1)}. " + _NUM_RE.sub("", raw, count=1))
                elif out_lines and out_lines[-1].startswith(("- ",)) and not para:
                    out_lines[-1] += " " + ln  # wrapped list item
                else:
                    para.append(ln)
            if para:
                out_lines.append(" ".join(para))
            items.append((b["bbox"][1], b["bbox"][0], "\n".join(out_lines)))
        items.sort(key=lambda t: (round(t[0], 0), t[1]))
        page_md = "\n\n".join(i[2] for i in items)
        parts.append(page_md)
    return "\n\n---\n\n".join(p for p in parts if p.strip()).strip() + "\n"


def export_html(doc: fitz.Document, title: str) -> str:
    pages = []
    for page in doc:
        pages.append(page.get_text("html"))
    css = ("body{background:#e5e7eb;margin:0;padding:24px;font-family:sans-serif}"
           "body>div{background:#fff;margin:0 auto 24px auto;box-shadow:0 1px 4px rgba(0,0,0,.2);"
           "position:relative;overflow:hidden}")
    return (f"<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>{_html.escape(title)}</title>"
            f"<style>{css}</style></head><body>\n" + "\n".join(pages) + "\n</body></html>\n")


def export_images_zip(doc: fitz.Document, fmt: str, dpi: int) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for page in doc:
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            if fmt == "jpg":
                data = pix.tobytes("jpeg", jpg_quality=90)
            else:
                data = pix.tobytes("png")
            zf.writestr(f"page_{page.number + 1:04d}.{fmt}", data)
    return buf.getvalue()


def _all_tables(doc: fitz.Document) -> list[tuple[int, int, list[list]]]:
    out = []
    for page in doc:
        for ti, t in enumerate(_find_tables(page)):
            rows = t.extract()
            if rows and any(any(c for c in r) for r in rows):
                out.append((page.number, ti, rows))
    return out


def _coerce_cell(v):
    if v is None:
        return None
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s) and len(s) < 16 and not (len(s) > 1 and s.startswith("0")):
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        try:
            return float(s)
        except ValueError:
            pass
    return s


def export_xlsx(doc: fitz.Document) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    tables = _all_tables(doc)
    if not tables:
        raise HTTPException(status_code=422, detail="No tables detected in this document")
    wb = Workbook()
    wb.remove(wb.active)
    for pno, ti, rows in tables:
        ws = wb.create_sheet(title=f"Page {pno + 1} Table {ti + 1}"[:31])
        for r in rows:
            ws.append([_coerce_cell(c) for c in r])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv_zip(doc: fitz.Document) -> bytes:
    import csv

    tables = _all_tables(doc)
    if not tables:
        raise HTTPException(status_code=422, detail="No tables detected in this document")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for pno, ti, rows in tables:
            s = io.StringIO()
            csv.writer(s).writerows([["" if c is None else c for c in r] for r in rows])
            zf.writestr(f"page_{pno + 1}_table_{ti + 1}.csv", s.getvalue())
    return buf.getvalue()


def export_docx(pdf_path: Path) -> bytes:
    from pdf2docx import Converter

    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "out.docx"
        cv = Converter(str(pdf_path))
        try:
            cv.convert(str(out))
        finally:
            cv.close()
        return out.read_bytes()


ExportFormat = Literal["docx", "txt", "md", "html", "png", "jpg", "xlsx", "csv"]

_MEDIA = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "txt": ("text/plain; charset=utf-8", "txt"),
    "md": ("text/markdown; charset=utf-8", "md"),
    "html": ("text/html; charset=utf-8", "html"),
    "png": ("application/zip", "png.zip"),
    "jpg": ("application/zip", "jpg.zip"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("application/zip", "csv.zip"),
}


@router.get("/{doc_id}/export/{fmt}")
async def export_document(doc_id: str, fmt: ExportFormat, dpi: int = 150, filename: str = "document"):
    path = _doc_path(doc_id)
    dpi = max(36, min(600, dpi))
    base = re.sub(r"[^A-Za-z0-9._ -]", "_", os.path.splitext(os.path.basename(filename))[0])[:80] or "document"
    media, ext = _MEDIA[fmt]

    def work() -> bytes:
        with _lock_for(doc_id):
            if fmt == "docx":
                return export_docx(path)
            doc = fitz.open(str(path))
            try:
                if fmt == "txt":
                    return export_txt(doc)
                if fmt == "md":
                    return export_markdown(doc).encode("utf-8")
                if fmt == "html":
                    return export_html(doc, base).encode("utf-8")
                if fmt in ("png", "jpg"):
                    return export_images_zip(doc, fmt, dpi)
                if fmt == "xlsx":
                    return export_xlsx(doc)
                return export_csv_zip(doc)
            finally:
                doc.close()

    data = await asyncio.to_thread(work)
    return _download(data, media, f"{base}.{ext}")


# ═════════════════════════════════════════════════════════════════════════════
#  Create PDF from files
# ═════════════════════════════════════════════════════════════════════════════

_PAGE_SIZES = {"letter": fitz.paper_rect("letter"), "a4": fitz.paper_rect("a4")}
_IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
_TEXT_EXT = {".txt", ".text"}
_MD_EXT = {".md", ".markdown"}

_STORY_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.35; color: #111; }
h1 { font-size: 22pt; margin: 0 0 8pt 0; } h2 { font-size: 17pt; margin: 10pt 0 6pt 0; }
h3 { font-size: 14pt; margin: 8pt 0 4pt 0; } h4,h5,h6 { font-size: 12pt; }
p { margin: 0 0 6pt 0; } pre, code { font-family: monospace; font-size: 9.5pt; }
pre { background: #f3f4f6; padding: 6pt; }
table { border-collapse: collapse; } td, th { border: 1px solid #999; padding: 3pt 5pt; }
blockquote { margin-left: 18pt; color: #444; }
"""


def _inline_md(s: str) -> str:
    s = _html.escape(s, quote=False)
    s = re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
    s = re.sub(r"\*\*\*(.+?)\*\*\*", r"<b><i>\1</i></b>", s)
    s = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: f"<b>{m.group(1) or m.group(2)}</b>", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)|(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)",
               lambda m: f"<i>{m.group(1) or m.group(2)}</i>", s)
    s = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r'<a href="\2">\1</a>', s)
    return s


def markdown_to_html(md: str) -> str:
    """Small, dependency-free Markdown -> HTML (headings, lists, code, tables, quotes)."""
    out: list[str] = []
    lines = md.replace("\r\n", "\n").split("\n")
    i = 0
    para: list[str] = []

    def flush():
        if para:
            out.append("<p>" + _inline_md(" ".join(para)) + "</p>")
            para.clear()

    while i < len(lines):
        ln = lines[i]
        st = ln.strip()
        if st.startswith("```"):
            flush()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            out.append("<pre>" + _html.escape("\n".join(code)) + "</pre>")
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", st)
        if m:
            flush()
            n = len(m.group(1))
            out.append(f"<h{n}>{_inline_md(m.group(2).strip())}</h{n}>")
            i += 1; continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", st):
            flush(); out.append("<hr/>"); i += 1; continue
        if st.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-{2,}", lines[i + 1].strip()):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|?\s*:?-{2,}", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            html_rows = []
            for ri, r in enumerate(rows):
                tag = "th" if ri == 0 else "td"
                html_rows.append("<tr>" + "".join(f"<{tag}>{_inline_md(c)}</{tag}>" for c in r) + "</tr>")
            out.append("<table>" + "".join(html_rows) + "</table>")
            continue
        if re.match(r"^\s*([-*+]|\d+[.)])\s+", ln):
            flush()
            ordered = bool(re.match(r"^\s*\d+[.)]", ln))
            tag = "ol" if ordered else "ul"
            items = []
            while i < len(lines) and re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]):
                items.append("<li>" + _inline_md(re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", lines[i])) + "</li>")
                i += 1
            out.append(f"<{tag}>" + "".join(items) + f"</{tag}>")
            continue
        if st.startswith(">"):
            flush()
            q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip()); i += 1
            out.append("<blockquote>" + _inline_md(" ".join(q)) + "</blockquote>")
            continue
        if not st:
            flush(); i += 1; continue
        para.append(st)
        i += 1
    flush()
    return "\n".join(out)


def text_to_html(text: str) -> str:
    return "\n".join(f"<p>{_html.escape(p) if p.strip() else '&#160;'}</p>"
                     for p in text.replace("\r\n", "\n").split("\n"))


def html_to_pdf(html_body: str, page_rect: fitz.Rect) -> fitz.Document:
    story = fitz.Story(html=html_body, user_css=_STORY_CSS)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    where = page_rect + (54, 54, -54, -54)
    more = 1
    while more:
        dev = writer.begin_page(page_rect)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return fitz.open("pdf", buf.getvalue())


def docx_to_html(data: bytes) -> str:
    """Best-effort DOCX -> HTML (paragraph styles, bold/italic/underline, lists, tables)."""
    import docx  # python-docx (installed as a pdf2docx dependency)
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(data))
    out = []

    def runs_html(p) -> str:
        s = []
        for r in p.runs:
            t = _html.escape(r.text)
            if not t:
                continue
            if r.bold: t = f"<b>{t}</b>"
            if r.italic: t = f"<i>{t}</i>"
            if r.underline: t = f"<u>{t}</u>"
            s.append(t)
        return "".join(s)

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            style = (p.style.name if p.style is not None else "") or ""
            body = runs_html(p)
            m = re.match(r"Heading (\d)", style)
            if style == "Title":
                out.append(f"<h1>{body}</h1>")
            elif m:
                n = min(6, int(m.group(1)) + (0 if int(m.group(1)) > 0 else 1))
                out.append(f"<h{n}>{body}</h{n}>")
            elif "List" in style:
                out.append(f"<ul><li>{body}</li></ul>")
            else:
                out.append(f"<p>{body or '&#160;'}</p>")
        elif tag == "tbl":
            t = Table(child, d)
            rows = []
            for row in t.rows:
                rows.append("<tr>" + "".join(f"<td>{_html.escape(c.text)}</td>" for c in row.cells) + "</tr>")
            out.append("<table>" + "".join(rows) + "</table>")
    return "\n".join(out)


def _image_page(dst: fitz.Document, data: bytes, page_size: str, margin: float):
    try:
        img_doc = fitz.open(stream=data)  # format sniffed from content
        pix_w, pix_h = img_doc[0].rect.width, img_doc[0].rect.height
        img_doc.close()
    except Exception:
        # Formats MuPDF can't read directly (e.g. some WEBP) -> Pillow -> PNG
        from PIL import Image
        im = Image.open(io.BytesIO(data))
        im.load()
        if im.mode not in ("RGB", "L", "RGBA"):
            im = im.convert("RGB")
        b = io.BytesIO(); im.save(b, "PNG"); data = b.getvalue()
        pix_w, pix_h = im.size
    if page_size == "fit":
        page = dst.new_page(width=pix_w, height=pix_h)
        page.insert_image(page.rect, stream=data)
        return
    base = _PAGE_SIZES[page_size]
    w, h = base.width, base.height
    if pix_w > pix_h:
        w, h = h, w  # landscape page for landscape images
    page = dst.new_page(width=w, height=h)
    page.insert_image(fitz.Rect(margin, margin, w - margin, h - margin), stream=data, keep_proportion=True)


def create_pdf_from_files(files: list[tuple[str, bytes]], page_size: str = "letter",
                          margin: float = 36) -> fitz.Document:
    if page_size not in ("letter", "a4", "fit"):
        raise HTTPException(status_code=400, detail="page_size must be letter, a4 or fit")
    out = fitz.open()
    text_rect = _PAGE_SIZES["a4" if page_size == "a4" else "letter"]
    skipped = []
    for name, data in files:
        ext = os.path.splitext(name.lower())[1]
        try:
            if ext in _IMG_EXT:
                _image_page(out, data, page_size, margin)
            elif ext == ".pdf":
                src = fitz.open("pdf", data)
                out.insert_pdf(src)
                src.close()
            elif ext in _MD_EXT:
                src = html_to_pdf(markdown_to_html(data.decode("utf-8", "replace")), text_rect)
                out.insert_pdf(src); src.close()
            elif ext in _TEXT_EXT:
                src = html_to_pdf(text_to_html(data.decode("utf-8", "replace")), text_rect)
                out.insert_pdf(src); src.close()
            elif ext == ".docx":
                src = html_to_pdf(docx_to_html(data), text_rect)
                out.insert_pdf(src); src.close()
            else:
                skipped.append(name)
        except HTTPException:
            raise
        except Exception as e:
            out.close()
            raise HTTPException(status_code=422, detail=f"Could not convert {name}: {e}")
    if skipped:
        out.close()
        raise HTTPException(status_code=400, detail=f"Unsupported file type(s): {', '.join(skipped)}")
    if len(out) == 0:
        out.close()
        raise HTTPException(status_code=400, detail="No pages were produced")
    return out


@router.post("/create")
async def create_pdf(
    files: list[UploadFile] = File(...),
    page_size: str = Form("letter"),
    filename: str = Form(""),
):
    """Combine images / text / markdown / docx / pdf files (in order) into a new document."""
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded")
    payload = []
    total = 0
    for f in files:
        data = await f.read()
        total += len(data)
        if total > MAX_CREATE_BYTES:
            raise HTTPException(status_code=413, detail="Uploaded files are too large")
        payload.append((os.path.basename(f.filename or "file"), data))

    doc = await asyncio.to_thread(create_pdf_from_files, payload, page_size)
    doc_id = str(uuid.uuid4())
    d = _upload_dir() / doc_id
    d.mkdir(parents=True)
    try:
        doc.set_metadata({**(doc.metadata or {}), "producer": "AI PDF Editor", "creator": "AI PDF Editor"})
        doc.save(str(d / "original.pdf"), garbage=3, deflate=True)
        (d / "annotations.json").write_text("{}")
        page_count = len(doc)
        metadata = doc.metadata
    except Exception:
        shutil.rmtree(d, ignore_errors=True)
        raise
    finally:
        doc.close()
    name = filename.strip() or (os.path.splitext(payload[0][0])[0] + ".pdf")
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    return {"id": doc_id, "filename": os.path.basename(name), "page_count": page_count, "metadata": metadata}


# ═════════════════════════════════════════════════════════════════════════════
#  Compress / optimize
# ═════════════════════════════════════════════════════════════════════════════

PRESETS = {
    "high": {"dpi": 300, "quality": 90, "subset_fonts": True},
    "balanced": {"dpi": 150, "quality": 75, "subset_fonts": True},
    "smallest": {"dpi": 96, "quality": 50, "subset_fonts": True},
}


class CompressRequest(BaseModel):
    preset: Literal["high", "balanced", "smallest"] = "balanced"
    target_dpi: Optional[int] = Field(None, ge=36, le=600)
    quality: Optional[int] = Field(None, ge=10, le=95)
    dry_run: bool = False


def _image_effective_dpi(doc: fitz.Document, xref: int, w: int, h: int) -> float:
    """Highest DPI at which this image is displayed anywhere in the document."""
    best = 0.0
    for page in doc:
        try:
            rects = page.get_image_rects(xref)
        except Exception:
            continue
        for r in rects:
            if r.width <= 0 or r.height <= 0:
                continue
            best = max(best, w / (r.width / 72), h / (r.height / 72))
    return best


def _rewrite_images(doc: fitz.Document, target_dpi: int, quality: int) -> dict:
    from PIL import Image

    seen = set()
    smask_xrefs = set()
    for page in doc:
        for img in page.get_images(full=True):
            if img[1]:
                smask_xrefs.add(img[1])
    stats = {"images_total": 0, "images_rewritten": 0, "image_bytes_before": 0, "image_bytes_after": 0}
    for page in doc:
        for img in page.get_images(full=True):
            xref, smask = img[0], img[1]
            if xref in seen or xref in smask_xrefs:
                continue
            seen.add(xref)
            stats["images_total"] += 1
            if smask:
                continue  # keep transparency intact
            try:
                if doc.xref_get_key(xref, "ImageMask")[1] == "true":
                    continue
                orig_len = len(doc.xref_stream_raw(xref) or b"")
                pix = fitz.Pixmap(doc, xref)
            except Exception:
                continue
            stats["image_bytes_before"] += orig_len
            if pix.alpha:
                pix = fitz.Pixmap(pix, 0)
            if pix.colorspace is None or pix.n not in (1, 3):
                try:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                except Exception:
                    stats["image_bytes_after"] += orig_len
                    continue
            mode = "L" if pix.n == 1 else "RGB"
            im = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
            eff = _image_effective_dpi(doc, xref, pix.width, pix.height)
            if eff > target_dpi * 1.05:
                f = target_dpi / eff
                new_size = (max(1, round(pix.width * f)), max(1, round(pix.height * f)))
                im = im.resize(new_size, Image.LANCZOS)
            b = io.BytesIO()
            im.save(b, "JPEG", quality=quality, optimize=True)
            new = b.getvalue()
            if len(new) < orig_len * 0.95:
                page.replace_image(xref, stream=new)
                stats["images_rewritten"] += 1
                stats["image_bytes_after"] += len(new)
            else:
                stats["image_bytes_after"] += orig_len
    return stats


def compress_pdf_bytes(src: bytes, preset: str, target_dpi: Optional[int] = None,
                       quality: Optional[int] = None) -> tuple[bytes, dict]:
    cfg = dict(PRESETS[preset])
    if target_dpi:
        cfg["dpi"] = target_dpi
    if quality:
        cfg["quality"] = quality
    doc = fitz.open("pdf", src)
    try:
        stats = _rewrite_images(doc, cfg["dpi"], cfg["quality"])
        fonts_subset = False
        if cfg["subset_fonts"]:
            try:
                doc.subset_fonts()
                fonts_subset = True
            except Exception:
                fonts_subset = False
        out = doc.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
                          clean=True, use_objstms=1)
    finally:
        doc.close()
    stats.update({"fonts_subset": fonts_subset, "target_dpi": cfg["dpi"], "quality": cfg["quality"]})
    return out, stats


@router.post("/{doc_id}/compress")
async def compress(doc_id: str, req: CompressRequest = Body(default_factory=CompressRequest)):
    path = _doc_path(doc_id)

    def work():
        with _lock_for(doc_id):
            before_bytes = path.read_bytes()
            out, stats = compress_pdf_bytes(before_bytes, req.preset, req.target_dpi, req.quality)
            before, after = len(before_bytes), len(out)
            applied = False
            if not req.dry_run and after < before:
                snapshot(doc_id, f"Compress ({req.preset})")
                tmp = str(path) + ".convert.tmp"
                Path(tmp).write_bytes(out)
                os.replace(tmp, str(path))
                applied = True
            return {
                "preset": req.preset,
                "before_bytes": before,
                "after_bytes": after if applied or req.dry_run else before,
                "optimized_bytes": after,
                "saved_bytes": max(0, before - after),
                "saved_percent": round(max(0, before - after) * 100 / before, 1) if before else 0.0,
                "applied": applied,
                "dry_run": req.dry_run,
                **stats,
            }

    return await asyncio.to_thread(work)
