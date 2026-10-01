"""convert2: Word -> PDF fidelity (images, header/footer, page setup, table
borders/shading, fonts, page breaks), optional Microsoft Word path, and the
reflowable/semantic HTML export.  Assertions are on the real output files."""

import io
import re
import uuid

import fitz
import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from PIL import Image

from backend import advanced_ops
from backend.features import convert


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(advanced_ops, "UPLOAD_DIR", tmp_path)
    return tmp_path


@pytest_asyncio.fixture
async def cl(store):
    app = FastAPI()
    app.include_router(convert.router)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        yield c


def make_docx() -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, RGBColor

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)  # A4
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin, sec.bottom_margin = Inches(1.2), Inches(1.0)
    sec.header.paragraphs[0].text = "ACME Confidential Header"
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    for kind, txt in (("begin", None), ("instr", " PAGE "), ("separate", None), ("text", "1"), ("end", None),
                      ("text", " of "), ("simple", "NUMPAGES")):
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); fp.add_run()._r.append(fc)
        elif kind == "instr":
            it = OxmlElement("w:instrText"); it.text = txt; fp.add_run()._r.append(it)
        elif kind == "simple":
            fs = OxmlElement("w:fldSimple"); fs.set(qn("w:instr"), " NUMPAGES "); fp._p.append(fs)
        else:
            fp.add_run(txt)
    d.add_heading("Docx Heading", level=1)
    p = d.add_paragraph("Plain then ")
    p.add_run("bolded").bold = True
    r = p.add_run(" serif red")
    r.font.name = "Times New Roman"
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    img = Image.new("RGB", (300, 150), (20, 160, 60))
    b = io.BytesIO(); img.save(b, "PNG"); b.seek(0)
    d.add_picture(b, width=Inches(2))
    d.add_paragraph("First bullet", style="List Bullet")
    d.add_paragraph("Second bullet", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "K1"
    t.cell(1, 1).text = "V2"
    tcpr = t.cell(0, 0)._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "FFFF00"); tcpr.append(shd)
    d.add_page_break()
    d.add_paragraph("Second page body")
    out = io.BytesIO(); d.save(out)
    return out.getvalue()


def test_docx_to_pdf_fidelity():
    doc = convert.docx_to_pdf(make_docx())
    assert doc.page_count == 2
    # page size + margins from the docx section (A4, 0.75in left margin)
    assert abs(doc[0].rect.width - 595.3) < 1 and abs(doc[0].rect.height - 841.9) < 1
    p0, p1 = doc[0], doc[1]
    head = p0.search_for("Docx Heading")[0]
    assert abs(head.x0 - 54) < 3 and head.y0 > 80  # inside left/top margin
    # header on every page in the top margin; footer with PAGE / NUMPAGES fields
    for i, p in enumerate(doc):
        h = p.search_for("ACME")  # (Story shapes "fi" as a ligature)
        assert h and h[0].y1 < 86.4, (i, h)
        f = p.search_for(f"Page {i + 1} of 2")
        assert f and f[0].y0 > p.rect.height - 72, (i, f)
    # inline image extracted from the package, sized from wp:extent (2in wide)
    imgs = p0.get_image_info()
    assert len(imgs) == 1 and abs(fitz.Rect(imgs[0]["bbox"]).width - 144) < 2
    # table: cell shading + borders
    fills = [d["fill"] for d in p0.get_drawings() if d.get("fill")]
    assert any(f and abs(f[0] - 1) < .01 and abs(f[1] - 1) < .01 and f[2] < .01 for f in fills), fills
    k1, v2 = p0.search_for("K1")[0], p0.search_for("V2")[0]
    border_rects = [fitz.Rect(d["rect"]) for d in p0.get_drawings()
                    if d.get("fill") and max(d["fill"]) < 0.1 and d["rect"].height < 2.5]
    assert any(r.y1 <= k1.y0 + 1 and r.x0 <= k1.x0 for r in border_rects)   # line above K1
    assert any(r.y0 >= v2.y1 - 1 for r in border_rects)                     # line under V2
    # fonts mapped: Times New Roman -> serif face, colour carried
    spans = [s for b in p0.get_text("dict")["blocks"] for l in b.get("lines", []) for s in l["spans"]]
    red = next(s for s in spans if "serif red" in s["text"])
    assert "Sans" not in red["font"] and red["color"] == 0xC00000
    assert any(s["text"].strip() == "bolded" and s["flags"] & 16 for s in spans)
    assert "•" in p0.get_text() and "First bullet" in p0.get_text()
    # hard page break honoured
    assert "Second page body" in p1.get_text() and "Second page body" not in p0.get_text()


async def test_create_from_docx_endpoint_and_word_engine(cl, store, monkeypatch):
    data = make_docx()
    r = await cl.post("/api/pdf/create", files=[("files", ("w.docx", data, "application/octet-stream"))])
    assert r.status_code == 200, r.text
    out = fitz.open(str(store / r.json()["id"] / "original.pdf"))
    assert out.page_count == 2 and out[0].get_image_info()
    out.close()

    caps = (await cl.get("/api/pdf/convert/capabilities")).json()
    assert caps["word"] == convert.word_available()

    monkeypatch.setattr(convert, "word_available", lambda: False)
    r = await cl.post("/api/pdf/create", data={"docx_engine": "word"},
                      files=[("files", ("w.docx", data, "application/octet-stream"))])
    assert r.status_code == 400 and "Word" in r.json()["detail"]

    seen = {}

    def fake_word(src, dst):
        seen["src"] = src.read_bytes()
        d = fitz.open(); d.new_page().insert_text((72, 72), "FROM WORD"); d.save(str(dst))

    monkeypatch.setattr(convert, "word_available", lambda: True)
    monkeypatch.setattr(convert, "_osascript_word_to_pdf", fake_word)
    r = await cl.post("/api/pdf/create", data={"docx_engine": "word"},
                      files=[("files", ("w.docx", data, "application/octet-stream"))])
    assert r.status_code == 200, r.text
    assert seen["src"] == data
    assert "FROM WORD" in fitz.open(str(store / r.json()["id"] / "original.pdf"))[0].get_text()


def _pdf_for_html() -> fitz.Document:
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_text((72, 80), "Annual Report", fontsize=24, fontname="hebo")
    p.insert_text((72, 120), "This is body text for the report.", fontsize=11)
    p.insert_text((72, 136), "It continues & has <markup>.", fontsize=11)
    p.insert_text((72, 170), "• First bullet item", fontsize=11)
    p.insert_text((72, 186), "• Second bullet item", fontsize=11)
    img = Image.new("RGB", (40, 20), (200, 0, 0)); b = io.BytesIO(); img.save(b, "PNG")
    p.insert_image(fitz.Rect(72, 200, 152, 240), stream=b.getvalue())
    x0, y0, cw, rh = 72, 280, 120, 24
    for r in range(4):
        p.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    for r, row in enumerate([["Item", "Qty", "Price"], ["Apples", "12", "3.50"], ["Pears", "7", "2.25"]]):
        for c, v in enumerate(row):
            p.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), v, fontsize=11)
    return doc


async def test_export_html_reflow_is_semantic(cl, store):
    doc_id = str(uuid.uuid4()); (store / doc_id).mkdir()
    _pdf_for_html().save(str(store / doc_id / "original.pdf"))
    r = await cl.get(f"/api/pdf/{doc_id}/export/html", params={"layout": "reflow"})
    assert r.status_code == 200
    h = r.text
    assert "position:absolute" not in h.replace(" ", "")
    assert re.search(r"<h1>\s*Annual Report\s*</h1>", h)
    assert "<p>This is body text for the report. It continues &amp; has &lt;markup&gt;.</p>" in h
    assert re.search(r"<ul>\s*<li>First bullet item</li>\s*<li>Second bullet item</li>\s*</ul>", h)
    assert re.search(r"<table><tr><th>Item</th><th>Qty</th><th>Price</th></tr><tr><td>Apples</td>", h)
    m = re.search(r'<img src="data:image/png;base64,([A-Za-z0-9+/=]+)"', h)
    assert m
    import base64
    im = Image.open(io.BytesIO(base64.b64decode(m.group(1))))
    assert im.size == (40, 20)
    # order: heading < paragraph < list < image < table
    idx = [h.index(s) for s in ("<h1>", "<p>This", "<ul>", "<img", "<table>")]
    assert idx == sorted(idx)
    # positioned stays the default
    pos = (await cl.get(f"/api/pdf/{doc_id}/export/html")).text
    assert "position:absolute" in pos.replace(" ", "")
