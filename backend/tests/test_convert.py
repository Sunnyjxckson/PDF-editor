"""Tests for backend/features/convert.py (OCR, export, create, compress).

The router is mounted on a private FastAPI app with UPLOAD_DIR pointed at a
tmp dir, so these tests do not depend on main.py wiring.
"""

import asyncio
import io
import json
import zipfile

import fitz
import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from PIL import Image, ImageDraw, ImageFont

from backend import advanced_ops
from backend.features import convert


# ─── fixtures / helpers ──────────────────────────────────────────────────────

@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(advanced_ops, "UPLOAD_DIR", tmp_path)
    return tmp_path


@pytest_asyncio.fixture
async def cl(store):
    app = FastAPI()
    app.include_router(convert.router)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        yield c


def put_pdf(store, doc: fitz.Document) -> str:
    import uuid
    doc_id = str(uuid.uuid4())
    d = store / doc_id
    d.mkdir()
    doc.save(str(d / "original.pdf"))
    doc.close()
    return doc_id


def open_doc(store, doc_id) -> fitz.Document:
    return fitz.open(str(store / doc_id / "original.pdf"))


def _font(size):
    for p in ("/System/Library/Fonts/Helvetica.ttc", "/Library/Fonts/Arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default(size=size)


def scanned_pdf(lines=("Quarterly Invoice Number 4821", "The quick brown fox jumps")) -> fitz.Document:
    """A page that contains ONLY an image of text (what a scanner produces)."""
    img = Image.new("RGB", (1700, 2200), "white")  # 200 dpi letter
    d = ImageDraw.Draw(img)
    f = _font(60)
    for i, line in enumerate(lines):
        d.text((150, 200 + i * 200), line, fill="black", font=f)
    b = io.BytesIO(); img.save(b, "PNG")
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_image(p.rect, stream=b.getvalue())
    return doc


def text_pdf() -> fitz.Document:
    """Heading, bold text, bullets, and a ruled 3x3 table."""
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_text((72, 80), "Annual Report", fontsize=24, fontname="hebo")
    p.insert_text((72, 120), "This is body text for the report.", fontsize=11, fontname="helv")
    p.insert_text((72, 136), "Revenue grew strongly this year.", fontsize=11, fontname="helv")
    p.insert_text((72, 160), "Important note", fontsize=11, fontname="hebo")
    p.insert_text((72, 190), "• First bullet item", fontsize=11, fontname="helv")
    p.insert_text((72, 206), "• Second bullet item", fontsize=11, fontname="helv")
    # ruled table
    x0, y0, cw, rh = 72, 260, 120, 24
    rows = [["Item", "Qty", "Price"], ["Apples", "12", "3.50"], ["Pears", "7", "2.25"]]
    for r in range(4):
        p.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        p.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            p.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), val, fontsize=11)
    p2 = doc.new_page(width=612, height=792)
    p2.insert_text((72, 80), "Second page text", fontsize=11)
    return doc


async def run_ocr_job(cl, doc_id, body=None):
    r = await cl.post(f"/api/pdf/{doc_id}/ocr", json=body or {})
    assert r.status_code == 200, r.text
    job_id = r.json()["job_id"]
    for _ in range(300):
        j = (await cl.get(f"/api/pdf/ocr/jobs/{job_id}")).json()
        if j["status"] in ("done", "error"):
            return j
        await asyncio.sleep(0.1)
    raise AssertionError("OCR job timed out")


# ─── OCR ──────────────────────────────────────────────────────────────────────

async def test_detect_scanned_vs_text(cl, store):
    doc = scanned_pdf()
    t = text_pdf()
    doc.insert_pdf(t)
    doc_id = put_pdf(store, doc)
    r = await cl.get(f"/api/pdf/{doc_id}/ocr/detect")
    assert r.status_code == 200
    data = r.json()
    assert data["scanned_pages"] == [0]
    assert data["needs_ocr"] == [0]
    assert data["pages"][0]["image_coverage"] > 0.9
    assert data["pages"][1]["text_chars"] > 50


async def test_ocr_languages(cl, store):
    r = await cl.get("/api/pdf/ocr/languages")
    assert "eng" in r.json()["languages"]


async def test_ocr_searchable_invisible_aligned_layer(cl, store):
    doc_id = put_pdf(store, scanned_pdf())
    before = open_doc(store, doc_id)
    assert before[0].get_text().strip() == ""
    pix_before = before[0].get_pixmap(dpi=72).samples
    before.close()

    job = await run_ocr_job(cl, doc_id, {"language": "eng"})
    assert job["status"] == "done", job
    assert job["progress"] == 1.0
    assert job["result"]["pages_processed"] == [0]
    assert job["result"]["words_added"] >= 8

    after = open_doc(store, doc_id)
    page = after[0]
    text = page.get_text()
    assert "Invoice" in text and "4821" in text and "brown" in text

    # Aligned: "Invoice" was drawn at x≈(150+ ~330)/200*72 in the 200dpi image,
    # line top at 200px -> 72pt.  The searchable hit must land on the scan.
    hits = page.search_for("Invoice")
    assert len(hits) == 1
    h = hits[0]
    assert 120 < h.x0 < 175 and 60 < h.y0 < 85 and 85 < h.y1 < 105, h

    # Invisible: every char is render mode 3 and the rendering is unchanged.
    traces = page.get_texttrace()
    assert traces and all(s["type"] == 3 for s in traces)
    assert page.get_pixmap(dpi=72).samples == pix_before
    after.close()

    # Undo snapshot was taken before the mutation.
    hist = json.loads((store / doc_id / "history.json").read_text())
    assert hist["versions"][-1]["operation"].startswith("OCR")

    # Running auto OCR again is a no-op (page already has an OCR layer).
    job2 = await run_ocr_job(cl, doc_id, {"language": "eng"})
    assert job2["result"]["total"] == 0


async def test_ocr_editable_produces_visible_text(cl, store):
    doc_id = put_pdf(store, scanned_pdf())
    job = await run_ocr_job(cl, doc_id, {"language": "eng", "mode": "editable"})
    assert job["status"] == "done", job
    page = open_doc(store, doc_id)[0]
    assert "Quarterly" in page.get_text()
    traces = page.get_texttrace()
    assert traces and all(s["type"] == 0 for s in traces)
    # ink colour sampled from the black scan text -> dark
    assert all(max(s["color"]) < 0.4 for s in traces)


async def test_ocr_explicit_pages_and_skip_text_pages(cl, store):
    doc = text_pdf()
    doc_id = put_pdf(store, doc)
    job = await run_ocr_job(cl, doc_id, {"language": "eng"})  # auto -> nothing scanned
    assert job["status"] == "done" and job["result"]["total"] == 0
    assert not (store / doc_id / "history.json").exists()
    job = await run_ocr_job(cl, doc_id, {"language": "eng", "pages": [9]})
    assert job["status"] == "error" and "Invalid page" in job["error"]


async def test_ocr_rejects_bad_language(cl, store):
    doc_id = put_pdf(store, scanned_pdf())
    r = await cl.post(f"/api/pdf/{doc_id}/ocr", json={"language": "klingon"})
    assert r.status_code == 400
    r = await cl.post(f"/api/pdf/{doc_id}/ocr", json={"language": "eng; rm -rf"})
    assert r.status_code == 400


async def test_invalid_doc_id(cl, store):
    assert (await cl.get("/api/pdf/../etc/ocr/detect")).status_code in (400, 404)
    assert (await cl.get("/api/pdf/not-a-uuid/export/txt")).status_code == 400
    assert (await cl.get("/api/pdf/00000000-0000-0000-0000-000000000000/export/txt")).status_code == 404


# ─── Export ───────────────────────────────────────────────────────────────────

async def test_export_txt(cl, store):
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/txt?filename=report.pdf")
    assert r.status_code == 200
    assert 'filename="report.txt"' in r.headers["content-disposition"]
    txt = r.content.decode()
    pages = txt.split("\f")
    assert len(pages) == 2
    assert "Annual Report" in pages[0] and "Apples" in pages[0]
    assert "Second page text" in pages[1]


async def test_export_markdown_structure(cl, store):
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/md")
    assert r.status_code == 200
    md = r.content.decode()
    assert "# Annual Report" in md
    assert "**Important note**" in md
    assert "- First bullet item" in md and "- Second bullet item" in md
    assert "| Item | Qty | Price |" in md
    assert "| Apples | 12 | 3.50 |" in md
    # table text not duplicated as loose paragraphs
    assert md.count("Apples") == 1
    assert "---" in md and "Second page text" in md


async def test_export_html(cl, store):
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/html")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html")
    h = r.content.decode()
    assert h.startswith("<!DOCTYPE html>")
    assert "Annual Report" in h and "Second page text" in h
    assert h.count('id="page0"') == 2  # one positioned div per page


@pytest.mark.parametrize("fmt,magic", [("png", b"\x89PNG"), ("jpg", b"\xff\xd8")])
async def test_export_images_zip(cl, store, fmt, magic):
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/{fmt}?dpi=72")
    assert r.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = sorted(zf.namelist())
    assert names == [f"page_0001.{fmt}", f"page_0002.{fmt}"]
    data = zf.read(names[0])
    assert data.startswith(magic)
    assert Image.open(io.BytesIO(data)).size == (612, 792)


async def test_export_xlsx_tables(cl, store):
    from openpyxl import load_workbook
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/xlsx")
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Page 1 Table 1"]
    rows = [list(row) for row in wb.active.iter_rows(values_only=True)]
    assert rows[0] == ["Item", "Qty", "Price"]
    assert rows[1] == ["Apples", 12, 3.5]
    assert rows[2] == ["Pears", 7, 2.25]


async def test_export_csv_and_no_tables(cl, store):
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/csv")
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    assert zf.namelist() == ["page_1_table_1.csv"]
    assert zf.read("page_1_table_1.csv").decode().splitlines()[1] == "Apples,12,3.50"

    doc = fitz.open(); doc.new_page().insert_text((72, 72), "no tables")
    doc_id2 = put_pdf(store, doc)
    assert (await cl.get(f"/api/pdf/{doc_id2}/export/xlsx")).status_code == 422


async def test_export_docx(cl, store):
    import docx
    doc_id = put_pdf(store, text_pdf())
    r = await cl.get(f"/api/pdf/{doc_id}/export/docx")
    assert r.status_code == 200
    d = docx.Document(io.BytesIO(r.content))
    body = "\n".join(p.text for p in d.paragraphs)
    assert "Annual Report" in body and "Second page text" in body
    assert len(d.tables) >= 1
    cells = [c.text.strip() for row in d.tables[0].rows for c in row.cells]
    assert "Apples" in cells and "3.50" in cells


async def test_export_bad_format(cl, store):
    doc_id = put_pdf(store, text_pdf())
    assert (await cl.get(f"/api/pdf/{doc_id}/export/exe")).status_code == 422


# ─── Create ───────────────────────────────────────────────────────────────────

def _png(size, color):
    b = io.BytesIO(); Image.new("RGB", size, color).save(b, "PNG"); return b.getvalue()


def _jpg(size, color):
    b = io.BytesIO(); Image.new("RGB", size, color).save(b, "JPEG"); return b.getvalue()


async def test_create_from_images(cl, store):
    files = [
        ("files", ("a.png", _png((800, 1000), "red"), "image/png")),
        ("files", ("b.jpg", _jpg((1200, 600), "blue"), "image/jpeg")),
    ]
    r = await cl.post("/api/pdf/create", files=files, data={"page_size": "letter"})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["page_count"] == 2 and j["filename"] == "a.pdf"
    doc = open_doc(store, j["id"])
    assert doc[0].rect.width < doc[0].rect.height  # portrait
    assert doc[1].rect.width > doc[1].rect.height  # landscape for wide image
    for i, page in enumerate(doc):
        imgs = page.get_images()
        assert len(imgs) == 1
        rect = page.get_image_rects(imgs[0][0])[0]
        assert rect in page.rect + (35, 35, -35, -35)
    # colours preserved: centre of page 1 is red, page 2 blue
    c0 = doc[0].get_pixmap(dpi=20).pixel(doc[0].get_pixmap(dpi=20).width // 2, doc[0].get_pixmap(dpi=20).height // 2)
    assert c0[0] > 200 and c0[2] < 50
    assert (store / j["id"] / "annotations.json").read_text() == "{}"


async def test_create_fit_page_size(cl, store):
    r = await cl.post("/api/pdf/create", files=[("files", ("x.png", _png((300, 200), "green"), "image/png"))],
                      data={"page_size": "fit"})
    doc = open_doc(store, r.json()["id"])
    # page = image at its own resolution; PNG without dpi metadata -> 96 dpi
    assert (doc[0].rect.width, doc[0].rect.height) == (225, 150)
    b = io.BytesIO(); Image.new("RGB", (300, 600), "red").save(b, "PNG", dpi=(300, 300))
    r = await cl.post("/api/pdf/create", files=[("files", ("y.png", b.getvalue(), "image/png"))],
                      data={"page_size": "fit"})
    doc = open_doc(store, r.json()["id"])
    assert (doc[0].rect.width, doc[0].rect.height) == (72, 144)


async def test_create_from_text_and_markdown(cl, store):
    md = b"# Project Plan\n\nSome **bold** words.\n\n- alpha\n- beta\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"
    long_txt = ("Line of plain text number %d\n" % 1).encode() + b"".join(
        f"Line of plain text number {i}\n".encode() for i in range(2, 120))
    files = [("files", ("plan.md", md, "text/markdown")), ("files", ("notes.txt", long_txt, "text/plain"))]
    r = await cl.post("/api/pdf/create", files=files, data={"filename": "Combined"})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["filename"] == "Combined.pdf"
    doc = open_doc(store, j["id"])
    assert j["page_count"] == len(doc) >= 3  # md page + txt spills to 2+ pages
    p0 = doc[0]
    assert "Project Plan" in p0.get_text() and "alpha" in p0.get_text()
    # heading bigger than body; bold run rendered with a bold font
    spans = [s for b in p0.get_text("dict")["blocks"] for l in b.get("lines", []) for s in l["spans"]]
    head = next(s for s in spans if "Project Plan" in s["text"])
    body = next(s for s in spans if "words" in s["text"])
    assert head["size"] > body["size"] * 1.5
    assert any(s["text"].strip() == "bold" and (s["flags"] & 16) for s in spans)
    all_text = "".join(p.get_text() for p in doc)
    assert "Line of plain text number 1\n" in all_text and "number 119" in all_text


async def test_create_from_docx(cl, store):
    import docx
    d = docx.Document()
    d.add_heading("Docx Heading", level=1)
    para = d.add_paragraph("Plain then ")
    para.add_run("bolded").bold = True
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "K1"; t.cell(1, 1).text = "V2"
    b = io.BytesIO(); d.save(b)
    r = await cl.post("/api/pdf/create", files=[("files", ("w.docx", b.getvalue(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    txt = open_doc(store, r.json()["id"])[0].get_text()
    assert "Docx Heading" in txt and "bolded" in txt and "K1" in txt and "V2" in txt


async def test_create_merges_pdf_and_rejects_unknown(cl, store):
    src = text_pdf(); pdf_bytes = src.tobytes(); src.close()
    r = await cl.post("/api/pdf/create", files=[
        ("files", ("in.pdf", pdf_bytes, "application/pdf")),
        ("files", ("z.png", _png((100, 100), "white"), "image/png"))])
    assert r.json()["page_count"] == 3
    r = await cl.post("/api/pdf/create", files=[("files", ("v.exe", b"MZ", "application/octet-stream"))])
    assert r.status_code == 400
    assert len([p for p in store.iterdir()]) == 1  # the failed create left no directory


# ─── Compress ─────────────────────────────────────────────────────────────────

def photo_pdf() -> fitz.Document:
    """A 2400x2400 noisy photo shown at 4in -> 600 dpi, plus real text."""
    noise = Image.effect_noise((2400, 2400), 40).convert("L")
    grad = Image.linear_gradient("L").resize((2400, 2400))
    img = Image.merge("RGB", (noise, grad, Image.blend(noise, grad, 0.5)))
    b = io.BytesIO(); img.save(b, "PNG")
    doc = fitz.open()
    p = doc.new_page(width=612, height=792)
    p.insert_image(fitz.Rect(72, 200, 360, 488), stream=b.getvalue())
    p.insert_text((72, 100), "Compress me but keep this text", fontsize=14)
    return doc


async def test_compress_balanced_downsamples_and_keeps_text(cl, store):
    doc_id = put_pdf(store, photo_pdf())
    path = store / doc_id / "original.pdf"
    size_before = path.stat().st_size

    r = await cl.post(f"/api/pdf/{doc_id}/compress", json={"preset": "balanced"})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["applied"] is True
    assert j["before_bytes"] == size_before
    assert j["after_bytes"] == path.stat().st_size < size_before * 0.5
    assert j["images_rewritten"] == 1

    doc = fitz.open(str(path))
    page = doc[0]
    assert "Compress me but keep this text" in page.get_text()
    xref = page.get_images()[0][0]
    info = doc.extract_image(xref)
    # 4in display at 150 dpi -> 600 px
    assert abs(info["width"] - 600) <= 2 and info["ext"] in ("jpeg", "jpg")
    assert page.get_image_rects(xref)[0] == fitz.Rect(72, 200, 360, 488)
    hist = json.loads((store / doc_id / "history.json").read_text())
    assert hist["versions"][-1]["operation"] == "Compress (balanced)"


async def test_compress_presets_ordering_and_dry_run(cl, store):
    sizes = {}
    for preset in ("high", "balanced", "smallest"):
        doc_id = put_pdf(store, photo_pdf())
        path = store / doc_id / "original.pdf"
        orig = path.read_bytes()
        r = await cl.post(f"/api/pdf/{doc_id}/compress", json={"preset": preset, "dry_run": True})
        j = r.json()
        assert j["applied"] is False and path.read_bytes() == orig  # dry run touched nothing
        sizes[preset] = j["optimized_bytes"]
    assert sizes["high"] > sizes["balanced"] > sizes["smallest"]


async def test_compress_never_grows_file(cl, store):
    doc = fitz.open(); doc.new_page().insert_text((72, 72), "tiny")
    doc_id = put_pdf(store, doc)
    path = store / doc_id / "original.pdf"
    orig = path.read_bytes()
    j = (await cl.post(f"/api/pdf/{doc_id}/compress", json={"preset": "smallest"})).json()
    if not j["applied"]:
        assert path.read_bytes() == orig and j["after_bytes"] == j["before_bytes"]
    else:
        assert path.stat().st_size < len(orig)
    assert "tiny" in fitz.open(str(path))[0].get_text()
